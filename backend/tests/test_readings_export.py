import csv
import datetime as dt
import io

from openpyxl import load_workbook

from app.models import Meter, Reading, Tenant
from tests.conftest import new_member


def _reading(db_session, meter, valor, created_at, **extra):
    reading = Reading(tenant_id=meter.tenant_id, meter_id=meter.id, valor=valor, created_at=created_at, **extra)
    db_session.add(reading)
    db_session.commit()
    return reading


def _export_csv(client, tenant, headers, year, month):
    resp = client.get(
        f"/api/t/{tenant.slug}/admin/readings/export",
        headers=headers,
        params={"year": year, "month": month, "formato": "csv"},
    )
    assert resp.status_code == 200
    return list(csv.DictReader(io.StringIO(resp.content.decode("utf-8-sig"))))


def test_export_has_one_row_per_meter_with_its_latest_reading_of_the_month(
    client, tenant, admin_headers, member, meter, gestor, db_session
):
    other = Meter(tenant_id=tenant.id, member_id=member.id, codigo="MED-002", tipo="agua", unidad="m3")
    db_session.add(other)
    db_session.commit()
    _reading(db_session, meter, 100, dt.datetime(2026, 8, 10, 15, 0))
    _reading(db_session, meter, 150, dt.datetime(2026, 9, 5, 15, 0))
    _reading(
        db_session, meter, 160, dt.datetime(2026, 9, 20, 15, 0),
        valor_anterior=100, consumo=60, gestor_id=gestor.id, anomala=True, observacion="Medidor cambiado",
    )
    _reading(db_session, other, 30, dt.datetime(2026, 9, 6, 15, 0))
    _reading(db_session, other, 45, dt.datetime(2026, 10, 1, 15, 0))

    rows = _export_csv(client, tenant, admin_headers, 2026, 9)

    assert [r["medidor"] for r in rows] == ["MED-001", "MED-002"]
    first = rows[0]
    assert first["numero_socio"] == member.numero_socio
    assert first["nombre"] == member.nombre
    assert float(first["lectura_actual"]) == 160
    assert float(first["lectura_anterior"]) == 100
    assert float(first["consumo"]) == 60
    assert first["fecha"] == "2026-09-20 12:00"
    assert first["gestor"] == gestor.nombre
    assert first["inusual"] == "si"
    assert first["observacion"] == "Medidor cambiado"
    assert float(rows[1]["lectura_actual"]) == 30
    assert rows[1]["lectura_anterior"] == ""
    assert rows[1]["inusual"] == "no"


def test_export_december_covers_the_whole_month(client, tenant, admin_headers, meter, db_session):
    _reading(db_session, meter, 10, dt.datetime(2026, 12, 31, 23, 0))
    _reading(db_session, meter, 20, dt.datetime(2027, 1, 1, 0, 0))

    rows = _export_csv(client, tenant, admin_headers, 2026, 12)

    assert [float(r["lectura_actual"]) for r in rows] == [10]


def test_export_leaves_out_other_tenants_readings(client, tenant, admin_headers, meter, db_session):
    foreign = Tenant(slug="otra-coop", name="Otra")
    db_session.add(foreign)
    db_session.commit()
    _reading(db_session, meter, 10, dt.datetime(2026, 9, 5))
    db_session.add(Reading(tenant_id=foreign.id, meter_id=meter.id, valor=999, created_at=dt.datetime(2026, 9, 6)))
    db_session.commit()

    rows = _export_csv(client, tenant, admin_headers, 2026, 9)

    assert [float(r["lectura_actual"]) for r in rows] == [10]


def test_export_xlsx_is_the_default_format(client, tenant, admin_headers, meter, db_session):
    _reading(db_session, meter, 77, dt.datetime(2026, 9, 5, 15, 0))

    resp = client.get(
        f"/api/t/{tenant.slug}/admin/readings/export", headers=admin_headers, params={"year": 2026, "month": 9}
    )

    assert resp.status_code == 200
    assert "lecturas-coopera-test-2026-09.xlsx" in resp.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(resp.content)).active
    header, row = list(sheet.iter_rows(values_only=True))
    assert header[2] == "medidor"
    assert row[2] == "MED-001"
    assert row[6] == 77


def test_export_with_no_readings_returns_only_the_header(client, tenant, admin_headers):
    assert _export_csv(client, tenant, admin_headers, 2026, 9) == []


def test_export_rejects_bad_period_and_requires_admin(client, tenant, admin_headers, gestor_headers):
    url = f"/api/t/{tenant.slug}/admin/readings/export"
    assert client.get(url, headers=admin_headers, params={"year": 2026, "month": 13}).status_code == 422
    assert client.get(url, params={"year": 2026, "month": 9}).status_code == 401
    assert client.get(url, headers=gestor_headers, params={"year": 2026, "month": 9}).status_code == 401


def test_export_includes_unread_meters_with_the_reading_left_empty(
    client, tenant, admin_headers, member, meter, db_session
):
    old = dt.datetime(2026, 1, 1)
    meter.created_at = old
    unread = Meter(
        tenant_id=tenant.id, member_id=member.id, codigo="MED-002", tipo="agua", unidad="m3", created_at=old
    )
    otros = [new_member(db_session, tenant, n).id for n in (1002, 1003, 1004)]
    never_read = Meter(
        tenant_id=tenant.id, member_id=otros[0], codigo="MED-003", unidad="kWh", lectura_inicial=500, created_at=old
    )
    inactive = Meter(tenant_id=tenant.id, member_id=otros[1], codigo="MED-004", activo=False, created_at=old)
    later = Meter(tenant_id=tenant.id, member_id=otros[2], codigo="MED-005", created_at=dt.datetime(2026, 10, 2))
    db_session.add_all([unread, never_read, inactive, later])
    db_session.commit()
    _reading(db_session, meter, 100, dt.datetime(2026, 9, 5, 15, 0))
    _reading(db_session, unread, 30, dt.datetime(2026, 8, 6, 15, 0))
    _reading(db_session, unread, 45, dt.datetime(2026, 10, 1, 15, 0))

    rows = _export_csv(client, tenant, admin_headers, 2026, 9)

    assert [r["medidor"] for r in rows] == ["MED-001", "MED-002", "MED-003"]
    assert rows[1]["lectura_actual"] == ""
    assert rows[1]["consumo"] == ""
    assert rows[1]["fecha"] == ""
    assert rows[1]["inusual"] == ""
    assert rows[1]["unidad"] == "m3"
    # Lo último que se sabía del medidor antes de ese mes.
    assert float(rows[1]["lectura_anterior"]) == 30
    assert float(rows[2]["lectura_anterior"]) == 500

    summary = client.get(
        f"/api/t/{tenant.slug}/admin/readings/export/summary",
        headers=admin_headers,
        params={"year": 2026, "month": 9},
    )
    assert summary.status_code == 200
    assert summary.json() == {"total": 3, "leidos": 1, "sin_leer": 2}


def test_export_keeps_the_reading_of_a_meter_deactivated_afterwards(
    client, tenant, admin_headers, meter, db_session
):
    _reading(db_session, meter, 100, dt.datetime(2026, 9, 5, 15, 0))
    meter.activo = False
    db_session.commit()

    rows = _export_csv(client, tenant, admin_headers, 2026, 9)

    assert [float(r["lectura_actual"]) for r in rows] == [100]
